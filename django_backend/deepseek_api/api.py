import json
import logging
import os
import re
import zipfile
from typing import Generator

from django.conf import settings
from django.db import connection
from django.http import StreamingHttpResponse
from ninja import File, NinjaAPI, Router
from ninja.files import UploadedFile as NinjaUploadedFile

from . import services
from .agents import LlmConfig, MultiAgentConfig, Orchestrator, OrchestratorInput
from .dashboard_stats import build_dashboard_stats
from .models import APIKey
from .schemas import ChatIn, ErrorResponse, HistoryOut, LoginIn, LoginOut
from .services import get_or_create_session, model_api_call

logger = logging.getLogger(__name__)

api = NinjaAPI(title="DeepSeek-R1:7B API", version="0.0.1")


def api_key_auth(request):
    """验证请求头中的 API Key。"""
    auth_header = request.headers.get("Authorization")
    if not auth_header:
        return None
    try:
        scheme, key = auth_header.split()
        if scheme.lower() != "bearer":
            return None
        api_key = APIKey.objects.get(key=key)
        if not api_key.is_valid():
            api_key.delete()
            return None
        return api_key
    except (ValueError, APIKey.DoesNotExist):
        return None


router = Router(auth=api_key_auth)


def _sse_line(payload: dict) -> str:
    return "data: " + json.dumps(payload, ensure_ascii=False) + "\n\n"


def _error_event(message: str, detail: dict | None = None) -> dict:
    payload = {
        "type": "error",
        "message": message,
        # 向后兼容旧前端字段
        "chunk": message,
    }
    if isinstance(detail, dict) and detail:
        payload["error_detail"] = detail
    return payload


def _sse_error_response(status_code: int, message: str) -> StreamingHttpResponse:
    return StreamingHttpResponse(
        _sse_line(_error_event(message)),
        status=status_code,
        content_type="text/event-stream",
    )


def _validate_office_archive(file_obj, filename: str) -> tuple[bool, str]:
    file_obj.seek(0)
    signature = file_obj.read(4)
    file_obj.seek(0)
    if signature != b"PK\x03\x04":
        return False, f"{filename} 文件结构不合法"

    with zipfile.ZipFile(file_obj) as archive:
        entries = archive.infolist()
        if len(entries) > settings.MAX_OFFICE_ARCHIVE_ENTRIES:
            return False, "Office 文件内条目过多，疑似压缩炸弹"

        total_uncompressed = sum(info.file_size for info in entries)
        if total_uncompressed > settings.MAX_OFFICE_UNCOMPRESSED_SIZE:
            return False, "Office 文件解压后体积过大，已拒绝"

    file_obj.seek(0)
    return True, ""


def clean_llm_reply(reply: str) -> str:
    """从模型原始回复中移除 <think>...</think> 标签块。"""
    return re.sub(r"<think>.*?</think>\s*", "", reply, flags=re.DOTALL).strip()


def _build_agent_cfg(base_llm: LlmConfig, agent_cfgs: dict, agent_id: str) -> LlmConfig:
    raw = agent_cfgs.get(agent_id) if isinstance(agent_cfgs, dict) else None
    if not isinstance(raw, dict):
        return base_llm
    return LlmConfig(
        provider=(raw.get("provider") or base_llm.provider or "ollama"),
        model=(raw.get("model") or base_llm.model),
        provider_api_key=(raw.get("provider_api_key") or base_llm.provider_api_key),
    )


def _build_multi_agent_config(base_llm: LlmConfig, agent_cfgs: dict) -> MultiAgentConfig:
    return MultiAgentConfig(
        rag=_build_agent_cfg(base_llm, agent_cfgs, "rag"),
        web=_build_agent_cfg(base_llm, agent_cfgs, "web"),
        synthesis=_build_agent_cfg(base_llm, agent_cfgs, "synthesis"),
    )


def _render_context(history_for_llm, user_input: str, final_reply: str) -> str:
    lines = []
    for msg in history_for_llm:
        if msg.get("role") == "user":
            lines.append(f"用户：{msg.get('content', '')}")
        elif msg.get("role") == "assistant":
            lines.append(f"回复：{msg.get('content', '')}")

    lines.append(f"用户：{user_input}")
    lines.append(f"回复：{final_reply}")
    return "\n".join(lines).strip()


@api.post("/login", response={200: LoginOut, 400: ErrorResponse, 403: ErrorResponse})
def login(request, data: LoginIn):
    username = data.username.strip()
    password = data.password.strip()
    if not username or not password:
        return 400, {"error": "用户名和密码不能为空"}
    if password != settings.AUTH_PASSWORD:
        return 403, {"error": "密码错误"}
    key = services.create_api_key(username)
    return {"api_key": key, "expiry": settings.TOKEN_EXPIRY_SECONDS}


@api.get("/health")
def health(request):
    return {
        "status": "ok",
        "service": "deepsoc-api",
    }


@api.get("/ready")
def ready(request):
    connection.ensure_connection()
    log_ready = services.log_system is not None
    status_code = 200 if log_ready else 503
    return status_code, {
        "status": "ready" if log_ready else "degraded",
        "db": "ok",
        "vector": "ok" if log_ready else "unavailable",
    }


@router.post("/chat")
def chat(request, data: ChatIn):
    if not request.auth:
        return _sse_error_response(401, "请先登录获取API Key")

    session_id = data.session_id.strip() or "默认对话"
    user_input = data.user_input.strip()
    if not user_input:
        return _sse_error_response(400, "请输入消息内容")

    user = request.auth
    session = get_or_create_session(session_id, user)

    use_db_search = data.use_db_search
    use_web_search = data.use_web_search
    selected_model = data.model_name
    selected_provider = (data.provider or "ollama").strip().lower()
    provider_api_key = data.provider_api_key
    web_search_api_key = data.web_search_api_key
    mode = (data.mode or "").strip().lower() or "single"

    logger.info(
        "搜索选项 - 数据库: %s, 联网: %s, provider: %s, model: %s",
        use_db_search,
        use_web_search,
        selected_provider,
        selected_model or "default",
    )

    if data.context and len(data.context) > 0:
        logger.info("使用前端提供的 context (会话: %s, 上下文长度: %s)", session_id, len(data.context))
        history_for_llm = data.context
        is_regeneration = False
    else:
        conversation_history = session.get_conversation_history()
        history_for_llm = conversation_history
        is_regeneration = False

        if (
            len(conversation_history) >= 2
            and conversation_history[-1].get("role") == "assistant"
            and conversation_history[-2].get("role") == "user"
            and conversation_history[-2].get("content") == user_input
        ):
            logger.info("检测到重新生成 (会话: %s)", session_id)
            history_for_llm = conversation_history[:-2]
            is_regeneration = True
        elif (
            len(conversation_history) >= 1
            and conversation_history[-1].get("role") == "user"
            and conversation_history[-1].get("content") == user_input
        ):
            logger.info("检测到对失败消息的重新生成 (会话: %s)", session_id)
            history_for_llm = conversation_history[:-1]
            is_regeneration = True

    def stream_generator() -> Generator[str, None, None]:
        full_clean_reply = ""

        try:
            if mode != "multi_agent":
                for raw_chunk in model_api_call(
                    user_input,
                    history_for_llm,
                    use_db_search,
                    use_web_search,
                    model_name=selected_model,
                    provider=selected_provider,
                    provider_api_key=provider_api_key,
                    web_search_api_key=web_search_api_key,
                ):
                    if not raw_chunk:
                        continue

                    if isinstance(raw_chunk, dict):
                        if raw_chunk.get("type") == "error":
                            msg = (raw_chunk.get("message") or raw_chunk.get("chunk") or "模型调用失败").strip()
                            detail = raw_chunk.get("error_detail")
                            yield _sse_line(_error_event(msg, detail if isinstance(detail, dict) else None))
                        else:
                            yield _sse_line(raw_chunk)
                        continue

                    yield _sse_line({"type": "content", "chunk": raw_chunk})
                    full_clean_reply += raw_chunk
            else:
                base_llm = LlmConfig(provider=selected_provider, model=selected_model, provider_api_key=provider_api_key)
                cfg = _build_multi_agent_config(base_llm, data.agent_configs or {})
                orch = Orchestrator()

                for evt in orch.run_stream(
                    OrchestratorInput(
                        query=user_input,
                        history=history_for_llm,
                        enable_rag=use_db_search,
                        enable_web=use_web_search,
                        web_search_api_key=web_search_api_key,
                    ),
                    cfg,
                ):
                    yield _sse_line(evt)
                    if evt.get("type") == "agent_chunk" and evt.get("agent_id") == "synthesis":
                        full_clean_reply += evt.get("content", "")

            final_save = full_clean_reply.strip()
            if not final_save:
                logger.warning("会话 %s 未收到有效模型输出，跳过写入", session_id)
                return

            if (data.context and len(data.context) > 0) or is_regeneration:
                session.context = _render_context(history_for_llm, user_input, final_save)
                session.save()
            else:
                session.update_context(user_input, final_save)

            logger.info("会话 %s 已更新 (用户: %s)", session_id, user.user)
        except Exception as e:
            logger.error("流式生成失败: %s", e)
            if mode == "multi_agent":
                detail = getattr(e, "detail", None)
                payload = {
                    "type": "agent_status",
                    "agent_id": "synthesis",
                    "status": "error",
                    "error": f"流处理失败: {e}",
                }
                if isinstance(detail, dict):
                    payload["error_detail"] = detail
                yield _sse_line(payload)
            else:
                yield _sse_line(_error_event(f"流处理失败: {e}"))
        finally:
            yield _sse_line({"type": "done"})

    response = StreamingHttpResponse(stream_generator(), content_type="text/event-stream")
    response["X-Accel-Buffering"] = "no"
    return response


@router.get("/history", response={200: HistoryOut})
def history(request, session_id: str = "默认对话"):
    processed_session_id = session_id.strip() or "默认对话"
    session = services.get_or_create_session(processed_session_id, request.auth)
    return {"history": session.context}


@router.delete("/history", response={200: dict})
def clear_history(request, session_id: str = "默认对话"):
    processed_session_id = session_id.strip() or "默认对话"
    session = services.get_or_create_session(processed_session_id, request.auth)
    session.clear_context()
    return {"message": "历史记录已清空"}


@router.get("/dashboard/stats", response={200: dict})
def get_dashboard_stats(request):
    """聚合 data/log 下真实 CSV 日志，返回大屏图表与拓扑数据。"""
    return build_dashboard_stats()


@router.post("/upload_file")
def upload_file(request, file: NinjaUploadedFile = File(...)):
    if not request.auth:
        return 401, {"error": "请先登录获取API Key"}

    raw_name = (file.name or "").strip()
    safe_name = os.path.basename(raw_name)
    if not safe_name:
        return 400, {"error": "文件名不能为空"}

    filename = safe_name.lower()
    _, ext = os.path.splitext(filename)
    allowed_exts = {".txt", ".docx", ".xlsx"}
    if ext not in allowed_exts:
        return 400, {"error": "不支持的文件类型，仅支持 .txt / .docx / .xlsx"}

    file_size = int(getattr(file, "size", 0) or 0)
    if file_size <= 0:
        return 400, {"error": "上传文件为空"}
    if file_size > settings.UPLOAD_MAX_BYTES:
        max_mb = settings.UPLOAD_MAX_BYTES // (1024 * 1024)
        return 400, {"error": f"文件过大，最大允许 {max_mb}MB"}

    try:
        if ext == ".txt":
            content_bytes = file.read()
            try:
                text = content_bytes.decode("utf-8")
            except Exception:
                text = content_bytes.decode("gbk", errors="ignore")
            return {"text": text}

        if ext == ".docx":
            try:
                from docx import Document
            except Exception:
                return 400, {"error": "缺少依赖：请安装 python-docx"}

            is_valid, reason = _validate_office_archive(file.file, safe_name)
            if not is_valid:
                return 400, {"error": reason}

            document = Document(file)
            paragraphs = [p.text for p in document.paragraphs if p.text]
            text = "\n".join(paragraphs)
            return {"text": text}

        if ext == ".xlsx":
            try:
                import openpyxl
            except Exception:
                return 400, {"error": "缺少依赖：请安装 openpyxl"}

            is_valid, reason = _validate_office_archive(file.file, safe_name)
            if not is_valid:
                return 400, {"error": reason}

            wb = openpyxl.load_workbook(file, data_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f"# 工作表: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = ["" if v is None else str(v) for v in row]
                    lines.append("\t".join(cells))
            text = "\n".join(lines)
            return {"text": text}

        return 400, {"error": "不支持的文件类型，仅支持 .txt / .docx / .xlsx"}

    except Exception as e:
        logger.error("文件解析失败: %s", e)
        return 400, {"error": f"文件解析失败: {e}"}


api.add_router("", router)
